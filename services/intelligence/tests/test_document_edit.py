"""Demo `/document/edit` endpoint tests over the real HTTP stack.

Same recipe as test_document_extract.py: the actual ASGI app under uvicorn
in a daemon thread (port 7902 so suites never share fate), stdlib urllib
only — no TestClient (httpx is not in STACK.md).
"""

import json
import os
import threading
import time
import urllib.error
import urllib.request

import pytest
import uvicorn

from agento_intelligence.main import app

TOKEN = os.environ["AGENTO_INTELLIGENCE_TOKEN"]
HOST = "127.0.0.1"
PORT = 7902
BASE_URL = f"http://{HOST}:{PORT}"


def _post(path: str, payload: dict, token: str | None = TOKEN) -> tuple[int, dict]:
    request = urllib.request.Request(
        BASE_URL + path,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    if token is not None:
        request.add_header("X-Agento-Token", token)
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        return error.code, json.loads(error.read().decode("utf-8"))


def _make_docx(path, with_bold: bool = False):
    import docx

    document = docx.Document()
    paragraph = document.add_paragraph("The quick brown fox jumps.")
    if with_bold:
        run = paragraph.add_run(" Keep this bold tail.")
        run.bold = True
    document.add_paragraph("A second paragraph stays untouched.")
    document.save(str(path))


def test_edit_paragraph_round_trip(server, tmp_path):
    target = tmp_path / "memo.docx"
    _make_docx(target)
    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "brown fox", "replacement": "red fox"}]},
    )
    assert status == 200
    assert body["edits_applied"] == 1
    assert "brown fox" in body["before_excerpt"]
    assert "red fox" in body["after_excerpt"]

    import docx

    texts = [p.text for p in docx.Document(str(target)).paragraphs]
    assert any("red fox" in text for text in texts)
    assert not any("brown fox" in text for text in texts)
    assert any("stays untouched" in text for text in texts)


def test_edit_preserves_run_formatting(server, tmp_path):
    target = tmp_path / "styled.docx"
    _make_docx(target, with_bold=True)
    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "quick brown", "replacement": "slow green"}]},
    )
    assert status == 200

    import docx

    first = docx.Document(str(target)).paragraphs[0]
    assert "slow green" in first.text
    bolds = [run.bold for run in first.runs if "Keep this bold tail." in run.text]
    assert bolds and all(bolds)


def test_edit_table_cell(server, tmp_path):
    import docx

    target = tmp_path / "table.docx"
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "price: 10"
    table.rows[0].cells[1].text = "other"
    document.save(str(target))

    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "price: 10", "replacement": "price: 12"}]},
    )
    assert status == 200
    saved = docx.Document(str(target))
    assert saved.tables[0].rows[0].cells[0].text == "price: 12"
    assert saved.tables[0].rows[0].cells[1].text == "other"


def test_edit_zero_match_is_422(server, tmp_path):
    target = tmp_path / "memo.docx"
    _make_docx(target)
    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "no such words", "replacement": "x"}]},
    )
    assert status == 422
    assert "does not match" in body["detail"]


def test_edit_multi_match_is_422(server, tmp_path):
    import docx

    target = tmp_path / "dupes.docx"
    document = docx.Document()
    document.add_paragraph("repeat after me")
    document.add_paragraph("repeat after me")
    document.save(str(target))
    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "repeat after me", "replacement": "x"}]},
    )
    assert status == 422
    assert "more than one place" in body["detail"]


def test_edit_missing_file_is_404(server, tmp_path):
    status, _body = _post(
        "/document/edit",
        {"path": str(tmp_path / "nope.docx"), "edits": [{"anchor": "a", "replacement": "b"}]},
    )
    assert status == 404


def test_edit_non_docx_is_422(server, tmp_path):
    target = tmp_path / "notes.txt"
    target.write_text("plain text", encoding="utf-8")
    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "plain", "replacement": "fancy"}]},
    )
    assert status == 422


def test_edit_pptx_round_trip(server, tmp_path):
    from pptx import Presentation

    target = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Revenue is up"
    presentation.save(str(target))

    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "Revenue is up", "replacement": "Revenue is flat"}]},
    )
    assert status == 200
    assert body["edits_applied"] == 1
    assert "Revenue is up" in body["before_excerpt"]
    assert "Revenue is flat" in body["after_excerpt"]

    saved = Presentation(str(target))
    bodies = [
        shape.text_frame.paragraphs[0].text
        for slide in saved.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert any("Revenue is flat" in text for text in bodies)
    assert not any("Revenue is up" in text for text in bodies)
    assert any("Quarterly Review" in text for text in bodies)


def test_edit_pptx_zero_match_is_422(server, tmp_path):
    from pptx import Presentation

    target = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    presentation.save(str(target))

    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "no such words", "replacement": "x"}]},
    )
    assert status == 422
    assert "presentation" in body["detail"]


def test_edit_xlsx_round_trip(server, tmp_path):
    from openpyxl import Workbook, load_workbook

    target = tmp_path / "prices.xlsx"
    workbook = Workbook()
    prices = workbook.active
    prices.title = "Prices"
    prices.append(["item", "price"])
    prices.append(["apples", 10])
    workbook.save(str(target))
    workbook.close()

    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "apples", "replacement": "pears"}]},
    )
    assert status == 200
    assert body["edits_applied"] == 1
    assert 'Sheet "Prices" A2' in body["before_excerpt"]
    assert "pears" in body["after_excerpt"]

    saved = load_workbook(str(target))
    assert saved["Prices"]["A2"].value == "pears"
    assert saved["Prices"]["B1"].value == "price"
    saved.close()


def test_edit_xlsx_substring_is_422(server, tmp_path):
    from openpyxl import Workbook

    target = tmp_path / "prices.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["granny smith apples"])
    workbook.save(str(target))
    workbook.close()

    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "apples", "replacement": "pears"}]},
    )
    assert status == 422
    assert "whole" in body["detail"]


def test_edit_xlsx_multi_match_is_422(server, tmp_path):
    from openpyxl import Workbook

    target = tmp_path / "dupes.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "One"
    first.append(["same"])
    second = workbook.create_sheet("Two")
    second.append(["same"])
    workbook.save(str(target))
    workbook.close()

    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "same", "replacement": "x"}]},
    )
    assert status == 422
    assert "more than one cell" in body["detail"]


def test_edit_tokenless_is_403(server, tmp_path):
    target = tmp_path / "memo.docx"
    _make_docx(target)
    status, body = _post(
        "/document/edit",
        {"path": str(target), "edits": [{"anchor": "fox", "replacement": "wolf"}]},
        token=None,
    )
    assert status == 403
    assert body["detail"] == "Forbidden: a valid X-Agento-Token header is required."


@pytest.fixture(scope="module")
def server():
    uvicorn_server = uvicorn.Server(
        uvicorn.Config(app, host=HOST, port=PORT, log_level="warning")
    )
    thread = threading.Thread(target=uvicorn_server.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10.0
    while not uvicorn_server.started and time.monotonic() < deadline:
        time.sleep(0.05)
    if not uvicorn_server.started:
        uvicorn_server.should_exit = True
        thread.join(timeout=5.0)
        raise RuntimeError(f"uvicorn did not start listening on {BASE_URL} in 10s")
    yield uvicorn_server
    uvicorn_server.should_exit = True
    thread.join(timeout=5.0)
