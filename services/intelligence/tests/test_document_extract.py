"""MVP /document/extract endpoint tests over the real HTTP stack.

Same recipe as test_health_smoke.py: the actual ASGI app under uvicorn in a
daemon thread (port 7901 so the two suites never share fate), stdlib urllib
only — no TestClient (httpx is not in STACK.md).
"""

import json
import os
import threading
import time
import urllib.error
import urllib.request

import pytest
import uvicorn

from agento_intelligence.main import app

TOKEN = os.environ["AGENTO_INTELLIGENCE_TOKEN"]
HOST = "127.0.0.1"
PORT = 7901
BASE_URL = f"http://{HOST}:{PORT}"


def _post(path: str, payload: dict, token: str | None = TOKEN) -> tuple[int, dict]:
    request = urllib.request.Request(
        BASE_URL + path,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    if token is not None:
        request.add_header("X-Agento-Token", token)
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        return error.code, json.loads(error.read().decode("utf-8"))


@pytest.fixture(scope="module")
def server():
    uvicorn_server = uvicorn.Server(
        uvicorn.Config(app, host=HOST, port=PORT, log_level="warning")
    )
    thread = threading.Thread(target=uvicorn_server.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10.0
    while not uvicorn_server.started and time.monotonic() < deadline:
        time.sleep(0.05)
    if not uvicorn_server.started:
        uvicorn_server.should_exit = True
        thread.join(timeout=5.0)
        raise RuntimeError(f"uvicorn did not start listening on {BASE_URL} in 10s")
    yield uvicorn_server
    uvicorn_server.should_exit = True
    thread.join(timeout=5.0)


def test_extract_txt(server, tmp_path):
    target = tmp_path / "notes.txt"
    target.write_text("first line\nsecond line", encoding="utf-8")
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert body == {"text": "first line\nsecond line", "truncated": False}


def test_extract_txt_truncated_at_cap(server, tmp_path):
    from agento_intelligence.extract import MAX_EXTRACT_CHARS

    target = tmp_path / "big.md"
    target.write_text("a" * (MAX_EXTRACT_CHARS + 100), encoding="utf-8")
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert body["truncated"] is True
    assert len(body["text"]) == MAX_EXTRACT_CHARS


def _minimal_pdf(text: str) -> bytes:
    """A tiny but fully valid single-page PDF (correct xref + startxref)."""
    stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def test_extract_pdf(server, tmp_path):
    target = tmp_path / "hello.pdf"
    target.write_bytes(_minimal_pdf("Hello Agento"))
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert body["truncated"] is False
    assert "Hello Agento" in body["text"]


def test_extract_docx(server, tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("First paragraph from Agento")
    document.add_paragraph("Second paragraph follows")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell-a"
    table.rows[0].cells[1].text = "cell-b"
    target = tmp_path / "report.docx"
    document.save(str(target))

    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert body["truncated"] is False
    assert "First paragraph from Agento" in body["text"]
    assert "Second paragraph follows" in body["text"]
    assert "cell-a\tcell-b" in body["text"]


def test_extract_pptx(server, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Revenue is up"
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    graphic = table_slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    graphic.table.cell(0, 0).text = "item"
    graphic.table.cell(0, 1).text = "price"
    target = tmp_path / "deck.pptx"
    presentation.save(str(target))

    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert body["truncated"] is False
    assert "--- Slide 1 ---" in body["text"]
    assert "Quarterly Review" in body["text"]
    assert "Revenue is up" in body["text"]
    assert "--- Slide 2 ---" in body["text"]
    assert "item\tprice" in body["text"]
    assert body["text"].index("Quarterly Review") < body["text"].index("item\tprice")


def test_extract_xlsx(server, tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    prices = workbook.active
    prices.title = "Prices"
    prices.append(["item", "price"])
    prices.append(["apples", 10])
    prices.append([])
    meta = workbook.create_sheet("Meta")
    meta.append(["owner", "sarah"])
    target = tmp_path / "prices.xlsx"
    workbook.save(str(target))

    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert body["truncated"] is False
    assert "--- Sheet: Prices ---" in body["text"]
    assert "item\tprice" in body["text"]
    assert "apples\t10" in body["text"]
    assert "--- Sheet: Meta ---" in body["text"]
    assert "owner\tsarah" in body["text"]


def test_extract_legacy_office_is_422(server, tmp_path):
    target = tmp_path / "old.xls"
    target.write_bytes(b"not really excel")
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 422
    assert ".xlsx" in body["detail"]


def test_extract_missing_file(server, tmp_path):
    missing = tmp_path / "nope.pdf"
    status, body = _post("/document/extract", {"path": str(missing)})
    assert status == 404
    # 2026-09-17: the absolute workspace path must never ride the detail —
    # the card sentence shows it verbatim. Basename only.
    assert str(tmp_path) not in body["detail"]
    assert "nope.pdf" in body["detail"]


def test_extract_unsupported_suffix(server, tmp_path):
    target = tmp_path / "photo.xyz"
    target.write_text("irrelevant", encoding="utf-8")
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 422
    assert "pdf" in body["detail"]


def test_extract_tokenless(server, tmp_path):
    target = tmp_path / "notes.txt"
    target.write_text("secret-ish", encoding="utf-8")
    status, body = _post("/document/extract", {"path": str(target)}, token=None)
    assert status == 403
    assert body["detail"] == "Forbidden: a valid X-Agento-Token header is required."


def test_extract_oversize_source_is_413(server, tmp_path, monkeypatch):
    import agento_intelligence.extract as extract_module

    monkeypatch.setattr(extract_module, "MAX_SOURCE_BYTES", 10)
    target = tmp_path / "notes.txt"
    target.write_text("x" * 11, encoding="utf-8")
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 413
    assert "too large" in body["detail"]
