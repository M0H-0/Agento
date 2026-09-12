"""Demo `/document/create` endpoint tests over the real HTTP stack.

Same recipe as test_document_edit.py: the actual ASGI app under uvicorn
in a daemon thread (port 7903 so suites never share fate), stdlib urllib
only — no TestClient (httpx is not in STACK.md).
"""

import json
import os
import threading
import time
import urllib.error
import urllib.request

import pytest
import uvicorn

from agento_intelligence.main import app

TOKEN = os.environ["AGENTO_INTELLIGENCE_TOKEN"]
HOST = "127.0.0.1"
PORT = 7903
BASE_URL = f"http://{HOST}:{PORT}"


def _post(path: str, payload: dict, token: str | None = TOKEN) -> tuple[int, dict]:
    request = urllib.request.Request(
        BASE_URL + path,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    if token is not None:
        request.add_header("X-Agento-Token", token)
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        return error.code, json.loads(error.read().decode("utf-8"))


def test_create_pptx_round_trip(server, tmp_path):
    target = tmp_path / "deck.pptx"
    status, body = _post(
        "/document/create",
        {
            "path": str(target),
            "title": "Roadmap",
            "items": ["Q1\nShip the beta", "Q2\nScale up"],
        },
    )
    assert status == 200
    assert body["size_bytes"] > 0
    assert "Roadmap" in body["after_excerpt"]
    assert "Ship the beta" in body["after_excerpt"]

    from pptx import Presentation

    saved = Presentation(str(target))
    assert len(saved.slides) == 3
    assert saved.slides[0].shapes.title.text == "Roadmap"
    assert saved.slides[1].shapes.title.text == "Q1"
    assert "Ship the beta" in saved.slides[1].placeholders[1].text

    # The created deck reads back through the extract ladder.
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert "--- Slide 1 ---" in body["text"]
    assert "Roadmap" in body["text"]
    assert "Ship the beta" in body["text"]


def test_create_xlsx_round_trip(server, tmp_path):
    target = tmp_path / "prices.xlsx"
    status, body = _post(
        "/document/create",
        {"path": str(target), "title": "Prices", "items": ["apples", "pears"]},
    )
    assert status == 200
    assert body["size_bytes"] > 0
    assert "apples" in body["after_excerpt"]

    from openpyxl import load_workbook

    saved = load_workbook(str(target))
    assert saved.sheetnames == ["Prices"]
    assert [cell.value for cell in saved["Prices"]["A"]] == ["apples", "pears"]
    saved.close()

    # The created workbook reads back through the extract ladder.
    status, body = _post("/document/extract", {"path": str(target)})
    assert status == 200
    assert "--- Sheet: Prices ---" in body["text"]
    assert "apples" in body["text"]


def test_create_unsupported_suffix_is_422(server, tmp_path):
    status, body = _post(
        "/document/create",
        {"path": str(tmp_path / "notes.pdf"), "title": "x", "items": ["y"]},
    )
    assert status == 422
    assert "write_file" in body["detail"]


def test_create_too_many_items_is_422(server, tmp_path):
    status, _body = _post(
        "/document/create",
        {"path": str(tmp_path / "big.pptx"), "title": "x", "items": ["y"] * 201},
    )
    assert status == 422


def test_create_tokenless_is_403(server, tmp_path):
    status, body = _post(
        "/document/create",
        {"path": str(tmp_path / "deck.pptx"), "title": "x", "items": ["y"]},
        token=None,
    )
    assert status == 403
    assert body["detail"] == "Forbidden: a valid X-Agento-Token header is required."


@pytest.fixture(scope="module")
def server():
    uvicorn_server = uvicorn.Server(
        uvicorn.Config(app, host=HOST, port=PORT, log_level="warning")
    )
    thread = threading.Thread(target=uvicorn_server.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10.0
    while not uvicorn_server.started and time.monotonic() < deadline:
        time.sleep(0.05)
    if not uvicorn_server.started:
        uvicorn_server.should_exit = True
        thread.join(timeout=5.0)
        raise RuntimeError(f"uvicorn did not start listening on {BASE_URL} in 10s")
    yield uvicorn_server
    uvicorn_server.should_exit = True
    thread.join(timeout=5.0)
